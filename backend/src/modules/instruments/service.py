from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

from fastapi import HTTPException, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession

from src.modules.instruments.models import (
    BookingStatus,
    Instrument,
    InstrumentBooking,
    InstrumentBookingApproval,
    InstrumentBookingRule,
    InstrumentStatusLog,
    InstrumentUsageRecord,
    ReviewStatus,
)
from src.modules.instruments.repository import InstrumentRepository
from src.modules.instruments.schemas import (
    ApprovalRequest,
    BookingListResponse,
    BookingRuleCreate,
    BookingRuleResponse,
    CalendarSlot,
    InstrumentBookingCreate,
    InstrumentBookingResponse,
    InstrumentCreate,
    InstrumentListResponse,
    InstrumentResponse,
    InstrumentUpdate,
    StatusChangeRequest,
    UsageRecordCreate,
    UsageRecordResponse,
)
from src.modules.users.models import User
from src.shared.notification_service import NotificationService


def _inst_response(i: Instrument) -> InstrumentResponse:
    return InstrumentResponse(
        id=i.id, code=i.code, name=i.name, model=i.model, manufacturer=i.manufacturer,
        serial_no=i.serial_no, asset_no=i.asset_no, category=i.category,
        purchase_date=i.purchase_date, purchase_price=float(i.purchase_price) if i.purchase_price else None,
        lab_id=i.lab_id, location=i.location, manager_id=i.manager_id, status=i.status,
        metadata=i.metadata_, lab_name=i.lab.name if i.lab else None,
        created_at=i.created_at, updated_at=i.updated_at,
    )


def _booking_response(b: InstrumentBooking) -> InstrumentBookingResponse:
    return InstrumentBookingResponse(
        id=b.id, instrument_id=b.instrument_id, instrument_name=b.instrument.name if b.instrument else None,
        user_id=b.user_id, start_time=b.start_time, end_time=b.end_time,
        purpose=b.purpose, project_name=b.project_name, status=b.status, created_at=b.created_at,
    )


class InstrumentService:
    def __init__(self, db: AsyncSession):
        self.repo = InstrumentRepository(db)
        self.notify = NotificationService(db)
        self.db = db

    async def _code(self) -> str:
        y = datetime.now(UTC).year
        return f"INS-{y}-{await self.repo.get_max_code_seq(y)+1:04d}"

    async def create(self, data: InstrumentCreate) -> InstrumentResponse:
        inst = Instrument(code=await self._code(), **{k: v for k, v in data.model_dump().items() if k != "metadata"}, metadata_=data.metadata or {})
        created = await self.repo.create_instrument(inst)
        await self.db.commit()
        return _inst_response(await self.repo.get_instrument(created.id))  # type: ignore

    async def get(self, iid: uuid.UUID) -> InstrumentResponse:
        inst = await self.repo.get_instrument(iid)
        if not inst:
            raise HTTPException(404, "仪器不存在")
        return _inst_response(inst)

    async def list(self, **kwargs) -> InstrumentListResponse:
        items, total = await self.repo.list_instruments(**kwargs)
        return InstrumentListResponse(items=[_inst_response(i) for i in items], total=total, page=kwargs.get("page", 1), page_size=kwargs.get("page_size", 20))

    async def update(self, iid: uuid.UUID, data: InstrumentUpdate, user: User) -> InstrumentResponse:
        inst = await self.repo.get_instrument(iid)
        if not inst:
            raise HTTPException(404, "仪器不存在")
        old_status = inst.status
        for k, v in data.model_dump(exclude_unset=True).items():
            if k == "metadata":
                inst.metadata_ = v
            else:
                setattr(inst, k, v)
        if data.status and data.status != old_status:
            await self.repo.add_status_log(InstrumentStatusLog(instrument_id=iid, old_status=old_status, new_status=data.status, reason=getattr(data, "reason", None), changed_by=user.id))
        await self.db.commit()
        return _inst_response(await self.repo.get_instrument(iid))  # type: ignore

    async def change_status(self, iid: uuid.UUID, data: StatusChangeRequest, user: User) -> InstrumentResponse:
        inst = await self.repo.get_instrument(iid)
        if not inst:
            raise HTTPException(404, "仪器不存在")
        old_status = inst.status
        inst.status = data.status
        await self.repo.add_status_log(InstrumentStatusLog(
            instrument_id=iid, old_status=old_status, new_status=data.status,
            reason=data.reason, changed_by=user.id,
        ))
        await self.db.commit()
        return _inst_response(await self.repo.get_instrument(iid))  # type: ignore

    async def set_rule(self, data: BookingRuleCreate) -> BookingRuleResponse:
        inst = await self.repo.get_instrument(data.instrument_id)
        if not inst:
            raise HTTPException(404, "仪器不存在")
        rule = await self.repo.get_rule(data.instrument_id)
        if rule:
            for k, v in data.model_dump().items():
                setattr(rule, k, v)
        else:
            rule = InstrumentBookingRule(**data.model_dump())
        saved = await self.repo.save_rule(rule)
        await self.db.commit()
        return BookingRuleResponse.model_validate(saved)

    async def get_rule(self, instrument_id: uuid.UUID) -> BookingRuleResponse:
        rule = await self.repo.get_rule(instrument_id)
        if not rule:
            raise HTTPException(404, "预约规则未配置")
        return BookingRuleResponse.model_validate(rule)

    async def create_booking(self, data: InstrumentBookingCreate, user: User) -> InstrumentBookingResponse:
        inst = await self.repo.get_instrument(data.instrument_id)
        if not inst or inst.status.value != "normal":
            raise HTTPException(400, "仪器不可用")
        if data.end_time <= data.start_time:
            raise HTTPException(400, "结束时间须晚于开始时间")
        if await self.repo.overlapping_bookings(data.instrument_id, data.start_time, data.end_time):
            raise HTTPException(409, "时段冲突")
        booking = InstrumentBooking(user_id=user.id, status=BookingStatus.PENDING, **data.model_dump())
        created = await self.repo.create_booking(booking)
        if inst.manager_id:
            await self.notify.send(inst.manager_id, "仪器预约待审批", f"{user.name} 申请预约 {inst.name}", "booking")
        await self.db.commit()
        refreshed = await self.repo.get_booking(created.id)
        return _booking_response(refreshed)  # type: ignore

    async def list_bookings(self, **kwargs) -> BookingListResponse:
        items, total = await self.repo.list_bookings(**kwargs)
        return BookingListResponse(items=[_booking_response(b) for b in items], total=total, page=kwargs.get("page", 1), page_size=kwargs.get("page_size", 20))

    async def approve_booking(self, bid: uuid.UUID, user: User, body: ApprovalRequest) -> InstrumentBookingResponse:
        booking = await self.repo.get_booking(bid)
        if not booking or booking.status != BookingStatus.PENDING:
            raise HTTPException(409, "不可审批")
        await self.repo.add_approval(InstrumentBookingApproval(booking_id=bid, approver_id=user.id, action="approve", comment=body.comment))
        booking.status = BookingStatus.APPROVED
        await self.notify.send(booking.user_id, "仪器预约已通过", "您的预约申请已批准", "booking")
        await self.db.commit()
        return _booking_response(await self.repo.get_booking(bid))  # type: ignore

    async def reject_booking(self, bid: uuid.UUID, user: User, body: ApprovalRequest) -> InstrumentBookingResponse:
        booking = await self.repo.get_booking(bid)
        if not booking or booking.status != BookingStatus.PENDING:
            raise HTTPException(409, "不可审批")
        await self.repo.add_approval(InstrumentBookingApproval(booking_id=bid, approver_id=user.id, action="reject", comment=body.comment))
        booking.status = BookingStatus.REJECTED
        await self.notify.send(booking.user_id, "仪器预约已驳回", body.comment or "预约被拒绝", "booking")
        await self.db.commit()
        return _booking_response(await self.repo.get_booking(bid))  # type: ignore

    async def calendar(self, instrument_id: uuid.UUID, from_time: datetime, to_time: datetime) -> list[CalendarSlot]:
        bookings, _ = await self.repo.list_bookings(instrument_id=instrument_id, from_time=from_time, to_time=to_time, page_size=200)
        return [CalendarSlot(start_time=b.start_time, end_time=b.end_time, status=b.status.value) for b in bookings]

    async def submit_usage(self, bid: uuid.UUID, data: UsageRecordCreate, user: User) -> UsageRecordResponse:
        booking = await self.repo.get_booking(bid)
        if not booking or booking.user_id != user.id:
            raise HTTPException(403, "无权操作")
        if await self.repo.get_usage(bid):
            raise HTTPException(409, "已提交使用记录")
        record = InstrumentUsageRecord(booking_id=bid, **data.model_dump())
        saved = await self.repo.save_usage(record)
        booking.status = BookingStatus.COMPLETED
        inst = await self.repo.get_instrument(booking.instrument_id)
        if inst and inst.manager_id:
            await self.notify.send(inst.manager_id, "仪器使用记录待审核", f"预约 {bid} 已提交使用记录", "usage")
        await self.db.commit()
        return UsageRecordResponse.model_validate(saved)

    async def review_usage(self, bid: uuid.UUID, approve: bool, user: User, comment: str | None = None) -> UsageRecordResponse:
        record = await self.repo.get_usage(bid)
        if not record:
            raise HTTPException(404, "使用记录不存在")
        record.review_status = ReviewStatus.APPROVED if approve else ReviewStatus.REJECTED
        record.reviewer_id = user.id
        record.reviewer_comment = comment
        await self.db.commit()
        await self.db.refresh(record)
        return UsageRecordResponse.model_validate(record)

    async def list_pending_usage(self, *, page: int = 1, page_size: int = 50):
        from src.modules.instruments.schemas import PendingInstrumentUsageItem, PendingInstrumentUsageListResponse

        records, total = await self.repo.list_pending_usage(page=page, page_size=page_size)
        items = [
            PendingInstrumentUsageItem(
                booking_id=r.booking_id,
                instrument_id=r.booking.instrument_id if r.booking else r.booking_id,
                instrument_name=r.booking.instrument.name if r.booking and r.booking.instrument else None,
                content=r.content,
                review_status=r.review_status,
                created_at=r.created_at,
            )
            for r in records
        ]
        return PendingInstrumentUsageListResponse(items=items, total=total)

    async def batch_review_usage(self, booking_ids: list[uuid.UUID], approve: bool, user: User, comment: str | None = None):
        from src.modules.instruments.schemas import BatchUsageReviewResponse

        processed = 0
        failed: list[uuid.UUID] = []
        for bid in booking_ids:
            record = await self.repo.get_usage(bid)
            if not record or record.review_status != ReviewStatus.PENDING:
                failed.append(bid)
                continue
            record.review_status = ReviewStatus.APPROVED if approve else ReviewStatus.REJECTED
            record.reviewer_id = user.id
            record.reviewer_comment = comment
            processed += 1
        await self.db.commit()
        return BatchUsageReviewResponse(processed=processed, failed=failed)

    async def export_instruments(self, fmt: str = "xlsx", **kwargs) -> tuple[bytes, str, str]:
        items, _ = await self.repo.list_instruments(page_size=5000, **kwargs)
        if fmt == "csv":
            lines = ["code,name,model,category,asset_no,status,lab_id,purchase_price"]
            for i in items:
                price = float(i.purchase_price) if i.purchase_price else ""
                lines.append(f"{i.code},{i.name},{i.model or ''},{i.category or ''},{i.asset_no or ''},{i.status.value},{i.lab_id or ''},{price}")
            return "\n".join(lines).encode("utf-8-sig"), "text/csv", "instruments.csv"
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(status_code=500, detail="openpyxl 未安装") from e
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "仪器设备"
        ws.append(["编号", "名称", "型号", "分类", "资产号", "状态", "实验室", "购置价格"])
        for i in items:
            ws.append([
                i.code, i.name, i.model, i.category, i.asset_no, i.status.value,
                i.lab.name if i.lab else "", float(i.purchase_price) if i.purchase_price else "",
            ])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "instruments.xlsx"

    async def import_instruments(self, file: UploadFile, user: User) -> dict:
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(status_code=500, detail="openpyxl 未安装") from e
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        imported = 0
        for row in rows:
            if not row or not row[0]:
                continue
            code = str(row[0]).strip()
            existing = await self.repo.get_instrument_by_code(code)
            if existing:
                if row[1]:
                    existing.name = str(row[1])
                if len(row) > 3 and row[3]:
                    existing.category = str(row[3])
            else:
                inst = Instrument(
                    code=code,
                    name=str(row[1]) if row[1] else code,
                    model=str(row[2]) if len(row) > 2 and row[2] else None,
                    category=str(row[3]) if len(row) > 3 and row[3] else None,
                    asset_no=str(row[4]) if len(row) > 4 and row[4] else None,
                    manager_id=user.id,
                )
                await self.repo.create_instrument(inst)
            imported += 1
        await self.db.commit()
        return {"imported": imported}
