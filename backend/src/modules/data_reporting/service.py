import io
import json
import uuid

from fastapi import HTTPException, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession

from src.modules.data_reporting.models import DataReportSubmission, DataReportTemplate, ReportSubmissionStatus
from src.modules.data_reporting.repository import DataReportingRepository
from src.modules.data_reporting.schemas import (
    AggregateStatsResponse,
    SubmissionCreate,
    SubmissionListResponse,
    SubmissionResponse,
    SubmissionUpdate,
    TemplateCreate,
    TemplateListResponse,
    TemplateResponse,
    TemplateUpdate,
)


def _template_response(t: DataReportTemplate) -> TemplateResponse:
    return TemplateResponse.model_validate(t)


def _submission_response(s: DataReportSubmission) -> SubmissionResponse:
    return SubmissionResponse(
        id=s.id,
        template_id=s.template_id,
        template_name=s.template.name if s.template else None,
        unit_name=s.unit_name,
        period=s.period,
        data=s.data,
        status=s.status,
        created_at=s.created_at,
        updated_at=s.updated_at,
    )


class DataReportingService:
    def __init__(self, db: AsyncSession):
        self.repo = DataReportingRepository(db)
        self.db = db

    async def create_template(self, data: TemplateCreate) -> TemplateResponse:
        if await self.repo.get_template_by_code(data.code):
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="模板代码已存在")
        template = DataReportTemplate(
            code=data.code,
            name=data.name,
            schema=data.schema_def,
            description=data.description,
        )
        created = await self.repo.create_template(template)
        await self.db.commit()
        return _template_response(created)

    async def get_template(self, template_id: uuid.UUID) -> TemplateResponse:
        template = await self.repo.get_template(template_id)
        if not template:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="模板不存在")
        return _template_response(template)

    async def list_templates(self, **kwargs) -> TemplateListResponse:
        items, total = await self.repo.list_templates(**kwargs)
        return TemplateListResponse(
            items=[_template_response(t) for t in items],
            total=total,
            page=kwargs.get("page", 1),
            page_size=kwargs.get("page_size", 20),
        )

    async def update_template(self, template_id: uuid.UUID, data: TemplateUpdate) -> TemplateResponse:
        template = await self.repo.get_template(template_id)
        if not template:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="模板不存在")
        if data.code and data.code != template.code:
            if await self.repo.get_template_by_code(data.code):
                raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="模板代码已存在")
        for k, v in data.model_dump(exclude_unset=True, by_alias=False).items():
            attr = "schema" if k == "schema_def" else k
            setattr(template, attr, v)
        await self.db.commit()
        await self.db.refresh(template)
        return _template_response(template)

    async def delete_template(self, template_id: uuid.UUID) -> None:
        template = await self.repo.get_template(template_id)
        if not template:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="模板不存在")
        await self.repo.delete_template(template)
        await self.db.commit()

    async def create_submission(self, data: SubmissionCreate) -> SubmissionResponse:
        template = await self.repo.get_template(data.template_id)
        if not template:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="模板不存在")
        submission = DataReportSubmission(**data.model_dump())
        created = await self.repo.create_submission(submission)
        await self.db.commit()
        refreshed = await self.repo.get_submission(created.id)
        return _submission_response(refreshed)  # type: ignore[arg-type]

    async def submit(self, submission_id: uuid.UUID) -> SubmissionResponse:
        submission = await self.repo.get_submission(submission_id)
        if not submission:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="报送记录不存在")
        if submission.status != ReportSubmissionStatus.DRAFT:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="仅草稿可提交")
        submission.status = ReportSubmissionStatus.SUBMITTED
        await self.db.commit()
        refreshed = await self.repo.get_submission(submission_id)
        return _submission_response(refreshed)  # type: ignore[arg-type]

    async def approve(self, submission_id: uuid.UUID) -> SubmissionResponse:
        submission = await self.repo.get_submission(submission_id)
        if not submission:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="报送记录不存在")
        if submission.status != ReportSubmissionStatus.SUBMITTED:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="仅已提交记录可审批")
        submission.status = ReportSubmissionStatus.APPROVED
        await self.db.commit()
        refreshed = await self.repo.get_submission(submission_id)
        return _submission_response(refreshed)  # type: ignore[arg-type]

    async def get_submission(self, submission_id: uuid.UUID) -> SubmissionResponse:
        submission = await self.repo.get_submission(submission_id)
        if not submission:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="报送记录不存在")
        return _submission_response(submission)

    async def list_submissions(self, **kwargs) -> SubmissionListResponse:
        items, total = await self.repo.list_submissions(**kwargs)
        return SubmissionListResponse(
            items=[_submission_response(s) for s in items],
            total=total,
            page=kwargs.get("page", 1),
            page_size=kwargs.get("page_size", 20),
        )

    async def update_submission(
        self, submission_id: uuid.UUID, data: SubmissionUpdate
    ) -> SubmissionResponse:
        submission = await self.repo.get_submission(submission_id)
        if not submission:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="报送记录不存在")
        if submission.status != ReportSubmissionStatus.DRAFT:
            raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="仅草稿可修改")
        for k, v in data.model_dump(exclude_unset=True).items():
            setattr(submission, k, v)
        await self.db.commit()
        refreshed = await self.repo.get_submission(submission_id)
        return _submission_response(refreshed)  # type: ignore[arg-type]

    async def aggregate_stats(self) -> AggregateStatsResponse:
        submissions = await self.repo.list_all_submissions()
        by_status: dict[str, int] = {}
        by_period: dict[str, int] = {}
        by_template: dict[str, int] = {}
        for s in submissions:
            by_status[s.status.value] = by_status.get(s.status.value, 0) + 1
            by_period[s.period] = by_period.get(s.period, 0) + 1
            tpl_name = s.template.name if s.template else str(s.template_id)
            by_template[tpl_name] = by_template.get(tpl_name, 0) + 1
        return AggregateStatsResponse(
            total_submissions=len(submissions),
            by_status=by_status,
            by_period=by_period,
            by_template=by_template,
        )

    async def export_submissions(self, template_id: uuid.UUID | None = None) -> bytes:
        items, _ = await self.repo.list_submissions(page_size=5000, template_id=template_id)
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(status_code=500, detail="openpyxl 未安装") from e
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "数据报送"
        ws.append(["单位", "周期", "模板", "状态", "数据"])
        for s in items:
            ws.append([
                s.unit_name,
                s.period,
                s.template.name if s.template else "",
                s.status.value,
                json.dumps(s.data, ensure_ascii=False),
            ])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    async def import_submissions(self, file: UploadFile) -> dict:
        content = await file.read()
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(status_code=500, detail="openpyxl 未安装") from e
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            template_name = str(row[2]) if len(row) > 2 and row[2] else None
            template = None
            if template_name:
                templates, _ = await self.repo.list_templates(page_size=100, keyword=template_name)
                template = templates[0] if templates else None
            if not template:
                continue
            data = {}
            if len(row) > 4 and row[4]:
                try:
                    data = json.loads(str(row[4]))
                except json.JSONDecodeError:
                    data = {"raw": str(row[4])}
            submission = DataReportSubmission(
                template_id=template.id,
                unit_name=str(row[0]),
                period=str(row[1]) if row[1] else "2025",
                data=data,
            )
            await self.repo.create_submission(submission)
            imported += 1
        await self.db.commit()
        return {"imported": imported}
