import io
import uuid
from datetime import UTC, datetime

from fastapi import HTTPException, UploadFile, status
from sqlalchemy.ext.asyncio import AsyncSession

from src.modules.labs.models import Lab, LabType, OpenStatus
from src.modules.labs.repository import LabRepository
from src.modules.labs.schemas import (
    LabCreate,
    LabImportResult,
    LabImportRow,
    LabListResponse,
    LabResponse,
    LabUpdate,
)
from src.modules.spaces.repository import SpaceRepository

LAB_TYPE_MAP = {
    "教学": LabType.TEACHING,
    "教学实验室": LabType.TEACHING,
    "teaching": LabType.TEACHING,
    "科研": LabType.RESEARCH,
    "科研实验室": LabType.RESEARCH,
    "research": LabType.RESEARCH,
    "综合": LabType.COMPREHENSIVE,
    "comprehensive": LabType.COMPREHENSIVE,
    "创新": LabType.INNOVATION,
    "innovation": LabType.INNOVATION,
    "实训": LabType.TRAINING,
    "training": LabType.TRAINING,
}

OPEN_STATUS_MAP = {
    "开放": OpenStatus.OPEN,
    "open": OpenStatus.OPEN,
    "关闭": OpenStatus.CLOSED,
    "closed": OpenStatus.CLOSED,
    "维护中": OpenStatus.MAINTENANCE,
    "维护": OpenStatus.MAINTENANCE,
    "maintenance": OpenStatus.MAINTENANCE,
}


def _lab_to_response(lab: Lab) -> LabResponse:
    building_name = floor_name = room_name = None
    if lab.room:
        room_name = lab.room.name
        if lab.room.floor:
            floor_name = lab.room.floor.name
            if lab.room.floor.building:
                building_name = lab.room.floor.building.name

    return LabResponse(
        id=lab.id,
        code=lab.code,
        name=lab.name,
        room_id=lab.room_id,
        location_detail=lab.location_detail,
        area_sqm=float(lab.area_sqm) if lab.area_sqm is not None else None,
        functional_zones=lab.functional_zones,
        capacity=lab.capacity,
        lab_type=lab.lab_type,
        manager_id=lab.manager_id,
        open_status=lab.open_status,
        inspection_status=lab.inspection_status,
        description=lab.description,
        metadata=lab.metadata_,
        created_at=lab.created_at,
        updated_at=lab.updated_at,
        building_name=building_name,
        floor_name=floor_name,
        room_name=room_name,
    )


class LabService:
    def __init__(self, db: AsyncSession):
        self.repo = LabRepository(db)
        self.space_repo = SpaceRepository(db)
        self.db = db

    async def _generate_code(self) -> str:
        year = datetime.now(UTC).year
        seq = await self.repo.get_max_code_seq(year) + 1
        return f"LAB-{year}-{seq:04d}"

    async def create(self, data: LabCreate) -> LabResponse:
        lab = Lab(
            code=await self._generate_code(),
            name=data.name,
            room_id=data.room_id,
            location_detail=data.location_detail,
            area_sqm=data.area_sqm,
            functional_zones=data.functional_zones or [],
            capacity=data.capacity,
            lab_type=data.lab_type,
            manager_id=data.manager_id,
            open_status=data.open_status,
            inspection_status=data.inspection_status,
            description=data.description,
            metadata_=data.metadata or {},
        )
        created = await self.repo.create(lab)
        await self.db.commit()
        refreshed = await self.repo.get_by_id(created.id)
        return _lab_to_response(refreshed)  # type: ignore[arg-type]

    async def get(self, lab_id: uuid.UUID) -> LabResponse:
        lab = await self.repo.get_by_id(lab_id)
        if not lab:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="实验室不存在")
        return _lab_to_response(lab)

    async def list(
        self,
        *,
        page: int = 1,
        page_size: int = 20,
        building_id: uuid.UUID | None = None,
        floor_id: uuid.UUID | None = None,
        lab_type: LabType | None = None,
        open_status: OpenStatus | None = None,
        open_statuses: list[OpenStatus] | None = None,
        keyword: str | None = None,
        manager_id: uuid.UUID | None = None,
    ) -> LabListResponse:
        items, total = await self.repo.list_labs(
            page=page,
            page_size=page_size,
            building_id=building_id,
            floor_id=floor_id,
            lab_type=lab_type,
            open_status=open_status,
            open_statuses=open_statuses,
            keyword=keyword,
            manager_id=manager_id,
        )
        return LabListResponse(
            items=[_lab_to_response(lab) for lab in items],
            total=total,
            page=page,
            page_size=page_size,
        )

    async def update(self, lab_id: uuid.UUID, data: LabUpdate) -> LabResponse:
        lab = await self.repo.get_by_id(lab_id)
        if not lab:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="实验室不存在")

        update_data = data.model_dump(exclude_unset=True)
        for field, value in update_data.items():
            if field == "metadata":
                setattr(lab, "metadata_", value)
            else:
                setattr(lab, field, value)

        await self.repo.update(lab)
        await self.db.commit()
        refreshed = await self.repo.get_by_id(lab_id)
        return _lab_to_response(refreshed)  # type: ignore[arg-type]

    async def delete(self, lab_id: uuid.UUID) -> None:
        lab = await self.repo.get_by_id(lab_id)
        if not lab:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="实验室不存在")
        await self.repo.soft_delete(lab)
        await self.db.commit()

    async def import_from_excel(self, file: UploadFile) -> LabImportResult:
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="openpyxl 未安装",
            ) from e

        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.active
        if ws is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="空 Excel 文件")

        success_count = 0
        errors: list[str] = []

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[1]:
                continue
            try:
                lab_type = None
                if row[8]:
                    lab_type = LAB_TYPE_MAP.get(str(row[8]).strip())
                open_status = OpenStatus.OPEN
                if row[9]:
                    open_status = OPEN_STATUS_MAP.get(str(row[9]).strip(), OpenStatus.OPEN)

                import_row = LabImportRow(
                    code=str(row[0]).strip() if row[0] else None,
                    name=str(row[1]).strip(),
                    building_name=str(row[2]).strip(),
                    floor_number=int(row[3]),
                    room_code=str(row[4]).strip() if row[4] else None,
                    location_detail=str(row[5]).strip() if row[5] else None,
                    area_sqm=float(row[6]) if row[6] else None,
                    capacity=int(row[7]) if row[7] else None,
                    lab_type=lab_type,
                    open_status=open_status,
                    description=str(row[10]).strip() if len(row) > 10 and row[10] else None,
                )

                room = await self.space_repo.find_room_by_location(
                    import_row.building_name,
                    import_row.floor_number,
                    import_row.room_code,
                )

                lab = Lab(
                    code=import_row.code or await self._generate_code(),
                    name=import_row.name,
                    room_id=room.id if room else None,
                    location_detail=import_row.location_detail,
                    area_sqm=import_row.area_sqm,
                    capacity=import_row.capacity,
                    lab_type=import_row.lab_type,
                    open_status=import_row.open_status or OpenStatus.OPEN,
                    description=import_row.description,
                )
                await self.repo.create(lab)
                success_count += 1
            except Exception as e:
                errors.append(f"第 {row_idx} 行: {e}")

        await self.db.commit()
        return LabImportResult(
            success_count=success_count,
            error_count=len(errors),
            errors=errors[:50],
        )

    async def export_to_excel(self) -> bytes:
        try:
            import openpyxl
        except ImportError as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="openpyxl 未安装",
            ) from e

        labs = await self.repo.list_all_for_export()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "实验室信息"
        headers = [
            "实验室编号", "实验室名称", "楼栋", "楼层", "房间",
            "具体位置", "面积", "容纳人数", "实验室类型", "开放状态", "备注",
        ]
        ws.append(headers)

        for lab in labs:
            building_name = floor_name = room_code = ""
            if lab.room:
                room_code = lab.room.code or lab.room.name
                if lab.room.floor:
                    floor_name = lab.room.floor.name
                    if lab.room.floor.building:
                        building_name = lab.room.floor.building.name

            ws.append([
                lab.code,
                lab.name,
                building_name,
                floor_name,
                room_code,
                lab.location_detail or "",
                float(lab.area_sqm) if lab.area_sqm else "",
                lab.capacity or "",
                lab.lab_type.value if lab.lab_type else "",
                lab.open_status.value,
                lab.description or "",
            ])

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
