#!/usr/bin/env python3
"""Export functional delivery report to Word (.docx) with embedded screenshots."""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from generate_delivery_report import ITEMS, STATUS_LABEL  # noqa: E402

SCREENSHOT_DIR = Path("/opt/cursor/artifacts/screenshots")
OUT_DOCX = ROOT / "docs" / "functional-delivery-report.docx"
ARTIFACT_DOCX = Path("/opt/cursor/artifacts/functional-delivery-report.docx")

MODULE_NAMES = {
    "1": "1. 实验室信息管理",
    "2": "2. 实验室设备管理",
    "3": "3. 实验室预约管理",
    "4": "4. 实验项目管理",
    "5": "5. 故障及问题上报",
    "6": "6. 数据填报管理",
    "7": "7. 数据对接与集成（部分排除）",
    "8": "8. 数据统计与分析",
    "9": "9. 系统基础功能",
    "10": "10. 智能体与接口",
    "11": "★11. 银校收费",
    "12": "★12. 实名制支付",
}


def _module_key(rid: str) -> str:
    if rid in ("11.1",):
        return "11"
    if rid in ("12.1",):
        return "12"
    return rid.split(".")[0]


def _set_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)


def _add_cover(doc: Document, *, done: int, excluded: int, total: int, in_scope: int) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("实验室管理系统\n功能交付报告")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = [
        f"报告日期：{datetime.now(timezone.utc).strftime('%Y年%m月%d日 %H:%M UTC')}",
        "系统版本：LabOS v0.3.3",
        "测试账号：admin/admin123",
        "",
        f"需求细项总数：{total} 条",
        f"本次交付范围：{in_scope} 条（排除集成对接 {excluded} 条）",
        f"范围内已完成：{done} 条（完成率 {100 * done // in_scope if in_scope else 0}%）",
    ]
    for line in lines:
        meta.add_run(line + "\n").font.size = Pt(12)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "说明：外部系统集成对接（模块 7、门禁/班牌/资产/人脸/数据中心等 production API）"
        "不在本次交付范围，报告中标记为「排除」。"
    ).font.size = Pt(10)
    note.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_page_break()


def _add_summary(doc: Document, test_summary: dict) -> None:
    doc.add_heading("一、交付结论", level=1)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("需求细项总数", str(len(ITEMS))),
        ("本次交付范围", f"{len(ITEMS) - sum(1 for i in ITEMS if i[2] == 'excluded')} 条"),
        ("范围内已完成", str(sum(1 for i in ITEMS if i[2] == "done"))),
        ("集成对接排除", str(sum(1 for i in ITEMS if i[2] == "excluded"))),
        (
            "范围内完成率",
            f"{sum(1 for i in ITEMS if i[2] == 'done')}/"
            f"{len(ITEMS) - sum(1 for i in ITEMS if i[2] == 'excluded')} = "
            f"{100 * sum(1 for i in ITEMS if i[2] == 'done') // max(1, len(ITEMS) - sum(1 for i in ITEMS if i[2] == 'excluded'))}%",
        ),
    ]
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v

    if test_summary:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run("自动化 API 测试：").bold = True
        p.add_run(
            f"PASS {test_summary.get('PASS', 0)} · "
            f"PARTIAL {test_summary.get('PARTIAL', 0)} · "
            f"FAIL {test_summary.get('FAIL', 0)}"
        )

    doc.add_page_break()


def _add_screenshot_index(doc: Document) -> None:
    doc.add_heading("二、截图索引", level=1)
    seen: set[str] = set()
    shots: list[str] = []
    for _, _, _, shot, _ in ITEMS:
        if shot and shot not in seen:
            seen.add(shot)
            shots.append(shot)

    doc.add_paragraph(f"本报告共引用 {len(shots)} 张功能截图（全量截图目录见 /opt/cursor/artifacts/screenshots/）。")
    table = doc.add_table(rows=len(shots) + 1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "文件名"
    table.rows[0].cells[1].text = "是否存在"
    for i, shot in enumerate(shots, start=1):
        table.rows[i].cells[0].text = shot
        exists = (SCREENSHOT_DIR / shot).exists()
        table.rows[i].cells[1].text = "✓" if exists else "（待补拍）"

    doc.add_page_break()


def _add_image(doc: Document, shot: str) -> None:
    path = SCREENSHOT_DIR / shot
    if not path.exists():
        doc.add_paragraph(f"【截图缺失：{shot}】").runs[0].font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        return
    doc.add_picture(str(path), width=Inches(6.2))
    cap = doc.add_paragraph(f"图：{shot}")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)


def _add_requirements(doc: Document) -> None:
    doc.add_heading("三、逐条功能说明（含截图）", level=1)
    current_module = ""

    for rid, title, status, shot, desc in ITEMS:
        mod_key = _module_key(rid)
        mod_name = MODULE_NAMES.get(mod_key, mod_key)
        if mod_name != current_module:
            current_module = mod_name
            doc.add_heading(mod_name, level=2)

        doc.add_heading(f"{rid} {title}", level=3)
        p = doc.add_paragraph()
        p.add_run("状态：").bold = True
        p.add_run(STATUS_LABEL[status])

        if shot:
            _add_image(doc, shot)

        for line in desc:
            doc.add_paragraph(line, style="List Bullet")

        doc.add_paragraph()


def _add_verification(doc: Document) -> None:
    doc.add_heading("四、验证方式", level=1)
    steps = [
        "cd backend && python3 -m scripts.reset_and_seed",
        "python3 -m uvicorn src.main:app --port 8000",
        "cd frontend && npm run dev",
        "浏览器访问 http://127.0.0.1:5173  账号 admin/admin123",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")


def main() -> None:
    test_path = ROOT / "docs" / "functional-test-report.json"
    test_summary = {}
    if test_path.exists():
        test_summary = json.loads(test_path.read_text()).get("summary", {})

    done = sum(1 for i in ITEMS if i[2] == "done")
    excluded = sum(1 for i in ITEMS if i[2] == "excluded")
    total = len(ITEMS)
    in_scope = total - excluded

    doc = Document()
    _set_default_style(doc)
    _add_cover(doc, done=done, excluded=excluded, total=total, in_scope=in_scope)
    _add_summary(doc, test_summary)
    _add_screenshot_index(doc)
    _add_requirements(doc)
    _add_verification(doc)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    ARTIFACT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    doc.save(str(ARTIFACT_DOCX))

    shot_count = len(list(SCREENSHOT_DIR.glob("*.png")))
    print(f"Word report saved:")
    print(f"  {OUT_DOCX}")
    print(f"  {ARTIFACT_DOCX}")
    print(f"Screenshots on disk: {shot_count}")


if __name__ == "__main__":
    main()
